#!/usr/bin/env python3
"""rt_build.py - resume-tailor builder (v1.1, 23 Oct 2026; council fixes).

One JSON source -> LaTeX (.tex) -> PDF, plus plain text (.txt) and optional Word (.docx),
then checks the output the way an ATS sees it.

Usage:
  python3 rt_build.py resume resume.json OUTDIR [--docx]
  python3 rt_build.py letter letter.json OUTDIR [--docx]

Resume layout is based on Jake Gutierrez's resume template (MIT License),
https://github.com/jakegut/resume . Needs: pdflatex (texlive-latex-extra), pdftotext, pdfinfo,
python-docx (only for --docx), pypdf (link check; skipped if missing).

Inline markup inside any text field: **bold** and [label](https://url).
Exit code: 0 = all hard checks passed, 1 = a hard check failed (see the printed report).
"""
import json, os, re, subprocess, sys, unicodedata

# ---------------------------------------------------------------- text helpers
LATEX_MAP = {'\\': r'\textbackslash{}', '&': r'\&', '%': r'\%', '$': r'\$', '#': r'\#',
             '_': r'{\texttt{\char95}}', '{': r'\{', '}': r'\}', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}',
             '|': r'$|$', '<': r'$<$', '>': r'$>$'}
UNICODE_OK = {'\u00a0': ' ', '–': '--', '’': "'", '‘': '`', '“': '``', '”': "''", '•': r'\textbullet{}',
              '×': r'$\times$', '…': r'\ldots{}', '★': '*', '·': r'\textperiodcentered{}', '→': r'$\rightarrow$'}
INLINE = re.compile(r'\*\*(.+?)\*\*|\[([^\]]+)\]\((https?://[^)\s]+|mailto:[^)\s]+|tel:[^)\s]+)\)')


def tex_plain(s):
    out = []
    for ch in s:
        if ch in LATEX_MAP:
            out.append(LATEX_MAP[ch])
        elif ch in UNICODE_OK:
            out.append(UNICODE_OK[ch])
        elif ord(ch) > 126 and unicodedata.category(ch)[0] not in 'L':
            raise ValueError(f'character {ch!r} (U+{ord(ch):04X}) not supported by pdfLaTeX template: "{s}"')
        else:
            out.append(ch)
    return ''.join(out)


def tex(s):
    """Inline markup -> LaTeX."""
    res, pos = [], 0
    for m in INLINE.finditer(s):
        res.append(tex_plain(s[pos:m.start()]))
        if m.group(1) is not None:
            res.append(r'\textbf{' + tex(m.group(1)) + '}')
        else:
            url = m.group(3).replace('%', r'\%').replace('#', r'\#')
            res.append(r'\href{' + url + '}{' + tex_plain(m.group(2)) + '}')
        pos = m.end()
    res.append(tex_plain(s[pos:]))
    return ''.join(res)


def plain(s, urls=True):
    """Inline markup -> plain text (links as 'label (url)' unless label already is the url)."""
    def rep(m):
        if m.group(1) is not None:
            return m.group(1)
        label, url = m.group(2), m.group(3)
        if not urls or url.startswith(('tel:', 'mailto:')):
            return label
        bare = re.sub(r'^(mailto:|https?://(www\.)?)', '', url).rstrip('/')
        return label if bare.lower() in label.lower() or label.lower() in bare.lower() else f'{label} ({url})'
    return INLINE.sub(rep, s)


def urls_in(s):
    return [m.group(3) for m in INLINE.finditer(s) if m.group(3)]


def slug(s):
    return re.sub(r'[^A-Za-z0-9]+', '_', s).strip('_')


# ---------------------------------------------------------------- LaTeX: resume
PREAMBLE = r"""%-------------------------
% Resume built by resume-tailor. Layout based on Jake Gutierrez's resume template
% (MIT License), https://github.com/jakegut/resume
% Paste into Overleaf: New Project > Blank Project > replace main.tex > Recompile (pdfLaTeX).
%-------------------------
\documentclass[__PAPER__paper,__FONTSIZE__pt]{article}
\usepackage{latexsym}
\usepackage[empty]{fullpage}
\usepackage{titlesec}
\usepackage[usenames,dvipsnames]{color}
\usepackage{enumitem}
\usepackage{hyperref}
\definecolor{linkblue}{rgb}{0,0.2,0.6}
\hypersetup{colorlinks=true,urlcolor=linkblue,linkcolor=linkblue}  % links shown in blue so recruiters see they are clickable
\usepackage{fancyhdr}
\usepackage[english]{babel}
\usepackage{tabularx}
\input{glyphtounicode}
\pdfgentounicode=1
\hyphenpenalty=10000
\exhyphenpenalty=10000
\hypersetup{pdftitle={__TITLE__},pdfauthor={__AUTHOR__}}
\pagestyle{fancy}
\fancyhf{}
\renewcommand{\headrulewidth}{0pt}
\renewcommand{\footrulewidth}{0pt}
__FOOTER__
\addtolength{\oddsidemargin}{-0.55in}
\addtolength{\evensidemargin}{-0.55in}
\addtolength{\textwidth}{1.1in}
\addtolength{\topmargin}{-0.6in}
\addtolength{\textheight}{1.15in}
\urlstyle{same}
\raggedbottom
\raggedright
\setlength{\tabcolsep}{0in}
\titleformat{\section}{\vspace{-4pt}\bfseries\raggedright\large}{}{0em}{}[\color{black}\titlerule \vspace{-5pt}]
\newcommand{\resumeItem}[1]{\item\small{{#1 \vspace{-2pt}}}}
\newcommand{\resumeSubheading}[4]{\vspace{-2pt}\item
  \begin{tabular*}{0.97\textwidth}[t]{l@{\extracolsep{\fill}}r}
    \textbf{#1} & #2 \\ \textit{\small#3} & \textit{\small #4} \\
  \end{tabular*}\vspace{-7pt}}
\newcommand{\resumeProjectHeading}[2]{\item
  \begin{tabular*}{0.97\textwidth}{l@{\extracolsep{\fill}}r}
    \small#1 & \small #2 \\
  \end{tabular*}\vspace{-7pt}}
\newcommand{\resumeInlineHeading}[4]{\vspace{-2pt}\item
  \textbf{#1}\ifx&#2&\else{} $|$ #2\fi \\ \textit{\small #3}\ifx&#4&\else{\small{} $|$ \textit{#4}}\fi \vspace{-1pt}}
\newcommand{\resumeSubHeadingListStart}{\begin{itemize}[leftmargin=0.15in, label={}]}
\newcommand{\resumeSubHeadingListEnd}{\end{itemize}}
\newcommand{\resumeItemListStart}{\begin{itemize}[label=$\bullet$]}
\newcommand{\resumeItemListEnd}{\end{itemize}\vspace{-5pt}}
\begin{document}
\thispagestyle{empty}
"""


def contact_parts(b):
    parts = []
    if b.get('location'):
        parts.append(b['location'])
    if b.get('phone'):
        tel = b.get('phone_e164') or re.sub(r'[^\d+]', '', b['phone'])
        parts.append(f"[{b['phone']}](tel:{tel})")
    if b.get('email'):
        parts.append(f"[{b['email']}](mailto:{b['email']})")
    for l in b.get('links', []):
        parts.append(f"[{l['label']}]({l['url']})")
    return parts


def build_resume_tex(d):
    m, b = d['meta'], d['basics']
    footer = ''
    if m.get('footer_name'):
        footer = r'\fancyfoot[C]{\small ' + tex_plain(b['name']) + r', page \thepage}'
    head = (PREAMBLE.replace('__PAPER__', m.get('paper', 'a4'))
            .replace('__FONTSIZE__', str(m.get('font_size', 11)))
            .replace('__TITLE__', tex_plain(f"{b['name']} - Resume"))
            .replace('__AUTHOR__', tex_plain(b['name']))
            .replace('__FOOTER__', footer))
    out = [head, r'\begin{center}', r'{\Huge \textbf{' + tex_plain(b['name']) + r'}} \\ \vspace{3pt}',
           r'\small ' + r' $|$ '.join(tex(p) for p in contact_parts(b))]
    for line in b.get('lines', []):
        out.append(r'\\ \vspace{1pt} \small \textbf{' + tex_plain(line['label']) + ':} ' + tex(line['text']))
    out.append(r'\end{center}')
    for s in d['sections']:
        out.append(r'\section{' + tex_plain(s['title']) + '}')
        t = s['type']
        if t == 'text':
            out.append(r'\small ' + tex(s['text']))
        elif t in ('summary', 'list'):
            if s.get('text'):
                out.append(r'\small ' + tex(s['text']))
            if s.get('bullets'):
                out.append(r'\begin{itemize}[leftmargin=0.15in, label=$\bullet$]')
                out += [r'\resumeItem{' + tex(x) + '}' for x in s['bullets']]
                out.append(r'\end{itemize}')
        elif t == 'entries':
            out.append(r'\resumeSubHeadingListStart')
            cmd = r'\resumeSubheading' if m.get('date_style', 'inline') == 'right' else r'\resumeInlineHeading'
            for e in s['items']:
                out.append(cmd + '{' + tex(e['title']) + '}{' + tex(e.get('dates', '')) + '}{' +
                           tex(e.get('org', '')) + '}{' + tex(e.get('location', '')) + '}')
                if e.get('note'):
                    out.append(r'\item[] \small ' + tex(e['note']))
                if e.get('bullets'):
                    out.append(r'\resumeItemListStart')
                    out += [r'\resumeItem{' + tex(x) + '}' for x in e['bullets']]
                    out.append(r'\resumeItemListEnd')
            out.append(r'\resumeSubHeadingListEnd')
        elif t == 'projects':
            out.append(r'\resumeSubHeadingListStart')
            for p in s['items']:
                left = r'\textbf{' + tex(p['name']) + '}'
                if p.get('meta'):
                    left += r' $|$ \emph{' + tex(p['meta']) + '}'
                if p.get('url'):
                    left += r' $|$ \href{' + p['url'].replace('%', r'\%').replace('#', r'\#') + '}{' + tex_plain(p.get('url_label', p['url'])) + '}'
                if m.get('date_style', 'inline') == 'right':
                    out.append(r'\resumeProjectHeading{' + left + '}{' + tex(p.get('dates', '')) + '}')
                else:
                    out.append(r'\item \small ' + left + (r' $|$ ' + tex(p['dates']) if p.get('dates') else '') + r' \vspace{-1pt}')
                if p.get('bullets'):
                    out.append(r'\resumeItemListStart')
                    out += [r'\resumeItem{' + tex(x) + '}' for x in p['bullets']]
                    out.append(r'\resumeItemListEnd')
            out.append(r'\resumeSubHeadingListEnd')
        elif t == 'skills':
            out.append(r'\begin{itemize}[leftmargin=0.15in, label={}]')
            out.append(r'\small{\item{')
            out.append(r' \\ '.join(r'\textbf{' + tex_plain(g['label']) + '}{: ' + tex(g['items']) + '}' for g in s['groups']))
            out.append(r'}}')
            out.append(r'\end{itemize}')
        else:
            raise ValueError('unknown section type ' + t)
    out.append(r'\end{document}')
    return '\n'.join(out) + '\n'


# ---------------------------------------------------------------- LaTeX: cover letter
def build_letter_tex(d):
    m, b = d['meta'], d['basics']
    lines = [r'% Cover letter built by resume-tailor. Paste into Overleaf like the resume (pdfLaTeX).',
             r'\documentclass[' + m.get('paper', 'a4') + 'paper,' + str(m.get('font_size', 11)) + 'pt]{article}',
             r'\usepackage[margin=0.9in]{geometry}', r'\usepackage{color}', r'\usepackage{hyperref}',
             r'\definecolor{linkblue}{rgb}{0,0.2,0.6}', r'\hypersetup{colorlinks=true,urlcolor=linkblue,linkcolor=linkblue}',
             r'\input{glyphtounicode}', r'\pdfgentounicode=1', r'\hyphenpenalty=10000', r'\exhyphenpenalty=10000',
             r'\hypersetup{pdftitle={' + tex_plain(b['name'] + ' - Cover letter') + '},pdfauthor={' + tex_plain(b['name']) + '}}',
             r'\setlength{\parindent}{0pt}', r'\setlength{\parskip}{8pt}', r'\pagestyle{empty}', r'\raggedright',
             r'\begin{document}', r'{\LARGE \textbf{' + tex_plain(b['name']) + r'}}\\[2pt]',
             r'\small ' + r' $|$ '.join(tex(p) for p in contact_parts(b)) + r'\normalsize', '', tex_plain(d['date'])]
    if d.get('recipient'):
        lines.append(r' \\ {}'.join(tex(x) for x in d['recipient']))
    if d.get('subject'):
        lines.append(r'\textbf{' + tex(d['subject']) + '}')
    lines.append(tex(d['greeting']))
    lines += [tex(p) for p in d['paragraphs']]
    lines.append(tex(d['closing']) + r'\\[18pt]' + tex_plain(d.get('signature', b['name'])))
    lines.append(r'\end{document}')
    return '\n\n'.join(lines) + '\n'


# ---------------------------------------------------------------- plain text
def build_resume_txt(d):
    b = d['basics']
    out = [b['name'], ' | '.join(plain(p) for p in contact_parts(b))]
    out += [f"{l['label']}: {plain(l['text'])}" for l in b.get('lines', [])]
    for s in d['sections']:
        out += ['', s['title'].upper()]
        t = s['type']
        if t == 'text':
            out.append(plain(s['text']))
        elif t in ('summary', 'list'):
            if s.get('text'):
                out.append(plain(s['text']))
            out += ['- ' + plain(x) for x in s.get('bullets', [])]
        elif t == 'entries':
            for e in s['items']:
                out.append(f"{plain(e['title'])} | {plain(e.get('org', ''))} | {plain(e.get('location', ''))} | {plain(e.get('dates', ''))}".replace(' |  |', ' |'))
                if e.get('note'):
                    out.append(plain(e['note']))
                out += ['- ' + plain(x) for x in e.get('bullets', [])]
        elif t == 'projects':
            for p in s['items']:
                bits = [plain(p['name'])] + [x for x in (plain(p.get('meta', '')), p.get('url', ''), plain(p.get('dates', ''))) if x]
                out.append(' | '.join(bits))
                out += ['- ' + plain(x) for x in p.get('bullets', [])]
        elif t == 'skills':
            out += [f"{g['label']}: {plain(g['items'])}" for g in s['groups']]
    return '\n'.join(out) + '\n'


def build_letter_txt(d):
    b = d['basics']
    out = [b['name'], ' | '.join(plain(p) for p in contact_parts(b)), '', d['date']]
    if d.get('recipient'):
        out += [plain(x) for x in d['recipient']]
    if d.get('subject'):
        out += ['', plain(d['subject'])]
    out += ['', plain(d['greeting']), '']
    for p in d['paragraphs']:
        out += [plain(p), '']
    out += [plain(d['closing']), d.get('signature', b['name'])]
    return '\n'.join(out) + '\n'


# ---------------------------------------------------------------- Word (.docx)
def build_docx(d, path, kind):
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    sec = doc.sections[0]
    if d['meta'].get('paper', 'a4') == 'a4':
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Inches(0.6))
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Calibri', Pt(10.5)
    usable = sec.page_width - sec.left_margin - sec.right_margin

    def add_runs(par, text):
        pos = 0
        for m in INLINE.finditer(text):
            if m.start() > pos:
                par.add_run(text[pos:m.start()])
            if m.group(1) is not None:
                par.add_run(m.group(1)).bold = True
            else:
                part = doc.part
                r_id = part.relate_to(m.group(3), 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), r_id)
                r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr'); u = OxmlElement('w:u'); u.set(qn('w:val'), 'single')
                rpr.append(u); r.append(rpr); t = OxmlElement('w:t'); t.text = m.group(2); r.append(t); h.append(r)
                par._p.append(h)
            pos = m.end()
        if pos < len(text):
            par.add_run(text[pos:])

    def heading(title):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(11.5)
        pPr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
        for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', '000000')):
            bot.set(qn(k), v)
        bdr.append(bot); pPr.append(bdr)

    def left_right(left, right, bold=False, italic=False):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(plain(left)); r.bold, r.italic = bold, italic
        if right:
            r2 = p.add_run('\t' + plain(right)); r2.italic = italic
        return p

    def bullets(items):
        for x in items:
            p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(0)
            add_runs(p, x)

    b = d['basics']
    p = doc.add_paragraph(); r = p.add_run(b['name']); r.bold = True; r.font.size = Pt(20)
    p = doc.add_paragraph(); add_runs(p, ' | '.join(contact_parts(b)))
    if kind == 'resume':
        for l in b.get('lines', []):
            p = doc.add_paragraph(); r = p.add_run(l['label'] + ': '); r.bold = True; add_runs(p, l['text'])
        for s in d['sections']:
            heading(s['title']); t = s['type']
            if t == 'text':
                add_runs(doc.add_paragraph(), s['text'])
            elif t in ('summary', 'list'):
                if s.get('text'):
                    add_runs(doc.add_paragraph(), s['text'])
                bullets(s.get('bullets', []))
            elif t == 'entries':
                for e in s['items']:
                    left_right(e['title'], e.get('dates', ''), bold=True)
                    left_right(e.get('org', ''), e.get('location', ''), italic=True)
                    if e.get('note'):
                        add_runs(doc.add_paragraph(), e['note'])
                    bullets(e.get('bullets', []))
            elif t == 'projects':
                for pj in s['items']:
                    left = pj['name'] + (' | ' + pj['meta'] if pj.get('meta') else '') + (' | ' + pj['url'] if pj.get('url') else '')
                    left_right(left, pj.get('dates', ''), bold=True)
                    bullets(pj.get('bullets', []))
            elif t == 'skills':
                for g in s['groups']:
                    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
                    r = p.add_run(g['label'] + ': '); r.bold = True; add_runs(p, g['items'])
    else:
        doc.add_paragraph(d['date'])
        if d.get('recipient'):
            doc.add_paragraph('\n'.join(plain(x) for x in d['recipient']))
        if d.get('subject'):
            p = doc.add_paragraph(); p.add_run(plain(d['subject'])).bold = True
        add_runs(doc.add_paragraph(), d['greeting'])
        for para in d['paragraphs']:
            add_runs(doc.add_paragraph(), para)
        doc.add_paragraph(plain(d['closing']))
        doc.add_paragraph(d.get('signature', b['name']))
    doc.core_properties.author = b['name']
    doc.core_properties.title = f"{b['name']} - {'Resume' if kind == 'resume' else 'Cover letter'}"
    doc.save(path)


# ---------------------------------------------------------------- checks
DUTY = re.compile(r'\b(responsible for|tasked with|worked on|duties included)\b', re.I)
AI_WORDS = ['spearheaded', 'leveraged', 'leveraging', 'orchestrated', 'passionate', 'results-driven', 'synergy',
            'cutting-edge', 'seamless', 'seamlessly', 'transformative', 'showcase', 'showcasing', 'delve', 'fostered',
            'harnessed', 'proven track record', 'fast-paced', 'tapestry', 'meticulous']
US_UK = [('optimiz', 'optimis'), ('organiz', 'organis'), ('analyz', r'analys(e|ed|ing)\b'), ('behavior', 'behaviour'),
         ('modeling', 'modelling'), ('color', 'colour'), ('prioritiz', 'prioritis'), ('standardiz', 'standardis'),
         ('utiliz', 'utilis'), ('minimiz', 'minimis'), ('maximiz', 'maximis'), ('summariz', 'summaris'),
         ('customiz', 'customis')]
YEAR_CTX = re.compile(r'(?:\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?|\b(?:in|since|from|until|to|by|of)|[\u2013-])\s*$', re.I)
NUM = re.compile(r'(?<![A-Za-z+#\d.,])\d[\d,]*(?:\.\d+)?')


def all_texts(d, kind, claims_only=False):
    """(where, text) pairs of the prose that lint and truth checks run on.
    claims_only=True keeps only claim-bearing prose (summary, bullets, notes, letter paragraphs)."""
    out = []
    if kind == 'letter':
        if not claims_only:
            out.append(('subject', d.get('subject', '')))
        return out + [('para', p) for p in d['paragraphs']]
    if not claims_only:
        for l in d['basics'].get('lines', []):
            out.append(('header', l['text']))
    for s in d['sections']:
        if s['type'] == 'skills' and claims_only:
            continue
        if s.get('text') and not (claims_only and s['type'] == 'text'):
            out.append((s['title'], s['text']))
        out += [(s['title'], x) for x in s.get('bullets', [])]
        for e in s.get('items', []):
            out += [(s['title'], x) for x in e.get('bullets', [])]
            if e.get('note'):
                out.append((s['title'], e['note']))
        if not claims_only:
            for g in s.get('groups', []):
                out.append((s['title'], g['items']))
    return out


def lint(d, kind, report):
    m = d['meta']
    variant = m.get('spelling', 'en-AU')
    allowed = set(str(x).replace(',', '').strip() for x in m.get('allowed_numbers', []))
    check_numbers = 'allowed_numbers' in m
    raw = json.dumps(d, ensure_ascii=False)
    if '\u2014' in raw or '---' in raw:
        report['hard'].append('em dash found (his rule: none)')
    if '[ADD' in raw or 'YOU DECIDE' in raw:
        report['draft'] = True
        report['warn'].append('placeholders present: files are a DRAFT')
    if variant not in ('en-AU', 'en-GB', 'en-NZ', 'en-IE', 'en-HK', 'en-SG', 'en-IN', 'en-US'):
        report['warn'].append(f'spelling variant {variant} not checked: proofread by hand')
    if kind == 'letter' and norm(m['company']) not in norm(' '.join([d.get('subject', '')] + d.get('recipient', []) + d['paragraphs'])):
        report['hard'].append(f'letter never names "{m["company"]}": wrong company or a copied letter?')
    body = json.dumps({k: v for k, v in d.items() if k != 'meta'}, ensure_ascii=False).replace('**', '')
    for w in m.get('never_use', []):
        if re.search(r'(?<!\w)' + re.escape(w) + r'(?!\w)', body, 0 if w.isupper() else re.I):
            report['hard'].append(f'never-use item "{w}" found (career doc section 0)')
    if 'never_use' not in m:
        report['warn'].append('meta.never_use missing: never-use check skipped')
    for where, t in all_texts(d, kind):
        p = plain(t)
        if kind == 'resume' and re.match(r"\s*(I|I'm|I've|My|Me)\b", p):
            report['hard'].append(f'first person at start [{where}]: {p[:60]}')
        if DUTY.search(p):
            report['hard'].append(f'duty phrase [{where}]: {DUTY.search(p).group(0)}')
        low = p.lower()
        for w in AI_WORDS:
            if re.search(r'\b' + re.escape(w) + r'\b', low):
                report['warn'].append(f'AI-sounding word "{w}" [{where}]')
        for us, uk in US_UK:
            if variant in ('en-AU', 'en-GB', 'en-NZ', 'en-IE', 'en-HK', 'en-SG', 'en-IN') and re.search(us, low):
                report['warn'].append(f'US spelling "{us}..." in {variant} [{where}]')
            if variant == 'en-US' and re.search(uk, low):
                report['warn'].append(f'UK spelling "{uk}..." in en-US [{where}]')
    if check_numbers:
        for where, t in all_texts(d, kind, claims_only=True):
            pt = plain(t, urls=False)
            for mm in NUM.finditer(pt):
                n = mm.group(0).rstrip(',.')
                tok = n.replace(',', '')
                if (',' not in n and re.fullmatch(r'(19|20)\d\d', tok)
                        and (YEAR_CTX.search(pt[:mm.start()]) or re.match(r'\s*[\u2013-]\s*(19|20)\d\d', pt[mm.start() + 4:]))):
                    continue
                if tok not in allowed:
                    report['hard'].append(f'number "{n}" not in meta.allowed_numbers [{where}]: {plain(t)[:60]}')
    else:
        report['warn'].append('meta.allowed_numbers missing: number truth check skipped')
    if kind == 'resume':
        for s in d['sections']:
            for e in s.get('items', []):
                bs = [plain(x).split()[0].lower() for x in e.get('bullets', []) if plain(x).split()]
                for a, c in zip(bs, bs[1:]):
                    if a == c:
                        report['warn'].append(f'two bullets in a row start with "{a}" [{s["title"]}]')


def norm(s):
    """Letters and digits only, lower case: survives line breaks, hyphens and quote styles."""
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKC', s).lower())


def check_pdf(pdf, d, kind, report):
    m, b = d['meta'], d['basics']
    info = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    pages = int(re.search(r'Pages:\s+(\d+)', info).group(1))
    w, h = map(float, re.search(r'Page size:\s+([\d.]+) x ([\d.]+)', info).groups())
    report['pages'] = pages
    maxp = m.get('pages_max', 1 if kind == 'letter' else 2)
    if pages > maxp:
        report['hard'].append(f'{pages} pages, limit {maxp}')
    want = (595, 842) if m.get('paper', 'a4') == 'a4' else (612, 792)
    if abs(w - want[0]) > 3 or abs(h - want[1]) > 3:
        report['hard'].append(f'page size {w:.0f}x{h:.0f} is not {m.get("paper", "a4")}')
    size_mb = os.path.getsize(pdf) / 1e6
    report['size_mb'] = round(size_mb, 3)
    if size_mb > min(1.0, m.get('size_limit_mb', 1.0)):
        report['hard'].append(f'PDF is {size_mb:.2f} MB (limit {min(1.0, m.get("size_limit_mb", 1.0))} MB)')
    txt = subprocess.run(['pdftotext', '-enc', 'UTF-8', pdf, '-'], capture_output=True, text=True).stdout
    report['pdftotext_chars'] = len(txt)
    if re.search('[ﬀ-ﬆ]', txt):
        report['hard'].append('ligature glyphs (U+FB00-FB06) in the text layer')
    flat = norm(txt)
    for need in [b['name'], b.get('email', ''), b.get('phone', '')]:
        if need and norm(need) not in flat:
            report['hard'].append(f'"{need}" missing from the text layer')
    for where, t in all_texts(d, kind):
        frag = norm(plain(t, urls=False))[:40]  # the PDF shows link labels, not URLs
        if frag and frag not in flat:
            report['hard'].append(f'text not found in PDF text layer [{where}]: {frag}')
    if kind == 'resume':
        lines = [norm(x) for x in txt.splitlines() if x.strip()]
        for s in d['sections']:
            for e in s.get('items', []):
                if s['type'] != 'entries' or not e.get('dates'):
                    continue
                tkey, dkey = norm(plain(e['title']))[:20], norm(plain(e['dates']))
                tis = [i for i, x in enumerate(lines) if tkey in x]
                dis = [i for i, x in enumerate(lines) if dkey in x]
                if not any(abs(a - c) <= 2 for a in tis for c in dis):
                    report['warn'].append(f'date "{e["dates"]}" may not sit next to "{e["title"]}" in reading order')
    try:
        from pypdf import PdfReader
        found = set()
        for page in PdfReader(pdf).pages:
            for a in page.get('/Annots') or []:
                obj = a.get_object()
                act = obj.get('/A')
                act = act.get_object() if act is not None else None
                uri = act.get('/URI') if act is not None else None
                if uri:
                    found.add(str(uri))
        wanted = set()
        for _, t in all_texts(d, kind):
            wanted.update(urls_in(t))
        for p in contact_parts(b):
            wanted.update(urls_in(p))
        for s in d.get('sections', []):
            for e in s.get('items', []):
                if e.get('url'):
                    wanted.add(e['url'])
        missing = [u for u in wanted if u not in found]
        if missing:
            report['hard'].append('links not clickable in PDF: ' + ', '.join(missing))
        report['links_ok'] = len(wanted) - len(missing)
    except ImportError:
        report['warn'].append('pypdf not installed: link check skipped (pip install pypdf --break-system-packages)')
    return txt


def check_docx(path, d, kind, report):
    from docx import Document
    doc = Document(path)
    if doc.tables:
        report['hard'].append('DOCX contains tables')
    text = norm('\n'.join(p.text for p in doc.paragraphs))
    for where, t in all_texts(d, kind):
        frag = norm(plain(t, urls=False))[:40]  # the PDF shows link labels, not URLs
        if frag and frag not in text:
            report['hard'].append(f'text missing from DOCX [{where}]: {frag}')


# ---------------------------------------------------------------- main
def main():
    if len(sys.argv) < 4 or sys.argv[1] not in ('resume', 'letter'):
        print(__doc__)
        sys.exit(2)
    kind, src, outdir = sys.argv[1], sys.argv[2], sys.argv[3]
    want_docx = '--docx' in sys.argv
    d = json.load(open(src, encoding='utf-8'))
    m = d['meta']
    os.makedirs(outdir, exist_ok=True)
    report = {'kind': kind, 'hard': [], 'warn': [], 'draft': False, 'files': []}
    lint(d, kind, report)
    stem = f"{m.get('file_stem', 'Chaitanya_Vankhande')}_{'Resume' if kind == 'resume' else 'Cover_Letter'}_{slug(m['company'])}"
    if report['draft']:
        stem += '_DRAFT'
    try:
        tex_src = build_resume_tex(d) if kind == 'resume' else build_letter_tex(d)
    except ValueError as err:
        report['hard'].append(str(err))
        report['status'] = 'FAIL'
        print(json.dumps(report, indent=1, ensure_ascii=False))
        sys.exit(1)
    tex_path = os.path.join(outdir, stem + '.tex')
    open(tex_path, 'w', encoding='utf-8').write(tex_src)
    for _ in range(2):
        r = subprocess.run(['pdflatex', '-interaction=nonstopmode', '-halt-on-error', '-output-directory', outdir, tex_path],
                           capture_output=True, text=True)
    pdf = os.path.join(outdir, stem + '.pdf')
    if r.returncode != 0 or not os.path.exists(pdf):
        errs = [l for l in r.stdout.splitlines() if l.startswith('!')][:5]
        report['hard'].append('pdflatex failed: ' + ' | '.join(errs))
        report['status'] = 'FAIL'
        print(json.dumps(report, indent=1, ensure_ascii=False))
        sys.exit(1)
    log_path = os.path.join(outdir, stem + '.log')
    if os.path.exists(log_path):
        for mm in re.finditer(r'Overfull \\hbox \(([\d.]+)pt too wide\)', open(log_path, encoding='latin-1').read()):
            if float(mm.group(1)) > 3:
                report['hard'].append(f'text runs {mm.group(1)}pt past the margin (overfull line): shorten or split it')
    for ext in ('.aux', '.log', '.out'):
        p = os.path.join(outdir, stem + ext)
        if os.path.exists(p):
            os.remove(p)
    check_pdf(pdf, d, kind, report)
    txt_path = os.path.join(outdir, stem + '.txt')
    open(txt_path, 'w', encoding='utf-8').write(build_resume_txt(d) if kind == 'resume' else build_letter_txt(d))
    report['files'] = [pdf, tex_path, txt_path]
    if want_docx:
        dpath = os.path.join(outdir, stem + '.docx')
        build_docx(d, dpath, kind)
        check_docx(dpath, d, kind, report)
        report['files'].append(dpath)
    if slug(m['company']).lower() not in os.path.basename(pdf).lower():
        report['hard'].append('company missing from file name')
    report['status'] = 'FAIL' if report['hard'] else ('DRAFT' if report['draft'] else 'PASS')
    print(json.dumps(report, indent=1, ensure_ascii=False))
    sys.exit(1 if report['hard'] else 0)


if __name__ == '__main__':
    main()
